"""Read the confirmed workbook without changing it. Output is private bootstrap data."""
import argparse, datetime, hashlib, json, re
from pathlib import Path
import openpyxl

KEYS = ['sn','model','productionDate','hardwareVersion','firmwareVersion','nickname','color','wifiMac','bluetoothMac','status','warrantyEnd','customer','deliveryDate','osVersion','aiVersion','asrVersion','ttsVersion','motionVersion','visionVersion','emotionVersion','otaVersion','cloudVersion','主控板','舵机控制器','头部舵机','左臂舵机','右臂舵机','腰部舵机','左腿舵机','右腿舵机','麦克风阵列','扬声器','摄像头','电池','传感器模组','assemblyPerson','assemblyDate','snEntryPerson','snEntryDate','debugPerson','debugDate','debugResult','qcPerson','qcDate','qcReport','remarks']
def norm(value):
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime('%Y-%m-%d')
    return value

def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = wb['机器人信息管理']
    if sheet.cell(3, 1).value != 'SN码' or sheet.cell(3, 13).value != '交付日期':
        raise ValueError('模板表头与已确认版本不一致，停止导入以避免错列')
    robots = []
    for row in sheet.iter_rows(min_row=4):
        sn = row[0].value
        if not isinstance(sn, str) or not re.fullmatch(r'LBR-\d{2}(0[1-9]|1[0-2])-[SPMR]-\d{4}', sn):
            continue
        values = [norm(cell.value) for cell in row[:46]]
        fields = {KEYS[i]: str(v) if v is not None else None for i,v in enumerate(values) if i and not 22 <= i <= 34}
        modules = [{'slot': KEYS[i], 'sn': str(values[i])} for i in range(22,35) if values[i] is not None]
        robots.append({'sn': sn, 'row': row[0].row, 'fields': fields, 'modules': modules})
    if len({r['sn'] for r in robots}) != len(robots):
        raise ValueError('原表包含重复 SN')
    sheets = [{'name': ws.title, 'cells': [{'cell':c.coordinate, 'row':c.row, 'column':c.column, 'value':norm(c.value)} for row in ws for c in row if c.value is not None]} for ws in wb]
    report = {
        'robotCount':len(robots), 'detailedRobots':sum(bool(r['fields']['model']) for r in robots),
        'moduleCount':sum(len(r['modules']) for r in robots), 'sheetCount':len(sheets),
        'mapping': [{'column':i+1,'sourceHeader':sheet.cell(3,i+1).value,'target':KEYS[i]} for i in range(46)],
        'decisions':[
            '采用确认文档 LBR 前缀和三个角色；原表说明中 R 前缀及四角色内容仅作来源留存。',
            'M 列为交付日期，独立存储，不当作激活日期。第一台状态为已交付。',
            '仅 SN 列匹配完整格式的行导入主档，图例和说明不作为机器人。',
            '八条仅 SN 记录保留所有业务空白；十个 SN 均已建主档占用，下一个 2609-P 为 0011。',
            '硬件模组版本和安装时间原表未提供，版本留空；installed_at 表示系统登记时间。',
            '测试结果“通过”属于整机调试，不等同质检结果。质检人员、日期、报告原样导入，qcResult 留空。',
            '初始已交付状态按原表保留；后续出库/交付需要新增明确通过的质检记录。',
            '全部三个工作表的非空单元格及坐标保存到 source_sheet，可在来源页面核对。'
        ]
    }
    return {'filename':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'robots':robots,'sheets':sheets,'report':report}

if __name__ == '__main__':
    p=argparse.ArgumentParser();p.add_argument('workbook',type=Path);p.add_argument('accounts',type=Path);p.add_argument('output',type=Path);a=p.parse_args()
    payload={'accounts':json.loads(a.accounts.read_text(encoding='utf-8-sig')),'workbook':extract(a.workbook)}
    a.output.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(payload['workbook']['report'],ensure_ascii=False,indent=2))
